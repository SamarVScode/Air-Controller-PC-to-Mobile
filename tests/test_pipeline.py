"""
Comprehensive QA and Simulation Test Suite for Daily Task Alert Pipeline.
Validates the end-to-end data pipeline against C:\\Users\\User\\Desktop\\res\\DAILY TASK VIEW.xlsx.

Replicates exact Google Apps Script logic from:
- HeaderResolver.js
- DateUtils.js
- TaskFilterService.js
- TelegramService.js
- Config.js
"""

import os
import sys
import re
import unittest
import hashlib
import datetime
import openpyxl

# Ensure UTF-8 output encoding for Windows consoles with emojis
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')


# ==============================================================================
# Pipeline Implementation (Replicating GAS Services Exactly)
# ==============================================================================

class HeaderResolver:
    """
    Implements 'The Header Map Law': Dynamic header mapping with alias resolution.
    ZERO hardcoded column array indices.
    """
    @staticmethod
    def normalize(header):
        if header is None:
            return ''
        s = str(header).strip().lower()
        s = re.sub(r'[\s\-_]+', '_', s)
        s = re.sub(r'[^a-z0-9_]', '', s)
        return s

    @classmethod
    def create(cls, header_row):
        if not header_row or not isinstance(header_row, (list, tuple)):
            raise ValueError("[HeaderResolver] Header row must be a non-empty array.")

        header_map = {}
        raw_headers = []

        for i, raw in enumerate(header_row):
            raw_str = str(raw) if raw is not None else ''
            raw_headers.append(raw_str)
            norm = cls.normalize(raw)
            if norm and norm not in header_map:
                header_map[norm] = i

        class ResolverInstance:
            def __init__(self, h_map, raw_h):
                self.map = h_map
                self.raw_headers = raw_h

            def get_index(self, aliases):
                alias_list = aliases if isinstance(aliases, (list, tuple)) else [aliases]
                for alias in alias_list:
                    norm = HeaderResolver.normalize(alias)
                    if norm in self.map:
                        return self.map[norm]
                raise KeyError(
                    f"[HeaderResolver] Missing required column matching aliases: {alias_list}. "
                    f"Available headers: {self.raw_headers}"
                )

            def resolve_indices(self, field_definitions):
                resolved = {}
                missing = []
                for key, aliases in field_definitions.items():
                    try:
                        resolved[key] = self.get_index(aliases)
                    except KeyError:
                        missing.append(key)

                if missing:
                    raise KeyError(
                        f"[HeaderResolver] Failed to resolve required fields: {missing}. "
                        f"Available headers: {self.raw_headers}"
                    )
                return resolved

        return ResolverInstance(header_map, raw_headers)


class DateUtils:
    """
    Implements 'Fast Date Handling': Zero Utilities.formatDate overhead.
    Standardized formatting to YYYY-MM-DD HH:mm:ss (or YYYY-MM-DD).
    """
    @staticmethod
    def pad2(n):
        return f"{int(n):02d}"

    @classmethod
    def format_fast(cls, value):
        if value is None or value == '':
            return 'N/A'

        if isinstance(value, datetime.datetime):
            y = f"{value.year:04d}"
            m = cls.pad2(value.month)
            d = cls.pad2(value.day)
            h = cls.pad2(value.hour)
            min_ = cls.pad2(value.minute)
            s = cls.pad2(value.second)
            return f"{y}-{m}-{d} {h}:{min_}:{s}"

        if isinstance(value, datetime.date):
            y = f"{value.year:04d}"
            m = cls.pad2(value.month)
            d = cls.pad2(value.day)
            return f"{y}-{m}-{d} 00:00:00"

        if isinstance(value, (int, float)):
            if value > 1000:
                # Excel epoch is 1899-12-30 (accounting for 1900 leap year bug)
                try:
                    excel_epoch = datetime.datetime(1899, 12, 30)
                    dt = excel_epoch + datetime.timedelta(days=value)
                    return cls.format_fast(dt)
                except Exception:
                    return 'Invalid Date'

        val_str = str(value).strip()
        if not val_str:
            return 'N/A'

        # Attempt string parse if looks like a date
        if len(val_str) >= 10 and any(c in val_str for c in ('-', '/', 'T')):
            formats_to_try = [
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d",
                "%Y/%m/%d %H:%M:%S",
                "%Y/%m/%d",
                "%d-%m-%Y %H:%M:%S",
                "%d-%m-%Y",
                "%d/%m/%Y %H:%M:%S",
                "%d/%m/%Y",
                "%Y-%m-%dT%H:%M:%S",
            ]
            for fmt in formats_to_try:
                try:
                    parsed = datetime.datetime.strptime(val_str, fmt)
                    return cls.format_fast(parsed)
                except ValueError:
                    continue

        return val_str


class TaskFilterService:
    """
    Handles spreadsheet tab resolution, zero-hardcoded column extraction,
    Hub MRZ filtering, and SHA-256 fingerprinting.
    """
    TAB_REGEX = re.compile(r'^tasks?\b', re.IGNORECASE)
    TARGET_HUB = 'MRZ'

    REQUIRED_FIELDS = {
        'tracking_number': ['final_tracking_number', 'tracking_number', 'tracking id', 'final_tracking_id', 'tracking'],
        'l5_name': ['l5_name', 'l5', 'l5 name', 'l5_user'],
        'leg': ['attribute', 'leg', 'task_leg', 'attributes'],
        'promise_date': ['promise_date', 'promise date', 'promised_date', 'promise'],
        'created_date': ['created_date', 'created date', 'creation_date', 'created'],
        'dc': ['dc', 'hub', 'delivery_center', 'dc_name']
    }

    @classmethod
    def resolve_task_sheet_name(cls, sheet_names):
        for name in sheet_names:
            if cls.TAB_REGEX.search(name.strip()):
                return name
        raise ValueError(
            f"[TaskFilterService] Could not find sheet matching pattern {cls.TAB_REGEX.pattern}. "
            f"Available sheets: {sheet_names}"
        )

    @classmethod
    def compute_fingerprint(cls, tasks):
        if not tasks:
            return 'EMPTY_SET'

        serialized = "\n".join(
            f"{t['trackingNumber']}|{t['l5Name']}|{t['leg']}|{t['promiseDate']}|{t['createdDate']}"
            for t in tasks
        )
        digest = hashlib.sha256(serialized.encode('utf-8')).hexdigest()
        return digest


class TelegramService:
    """
    Handles Telegram card formatting, HTML escaping, and <3800 char batch chunking.
    """
    MAX_MESSAGE_LENGTH = 3800

    @staticmethod
    def escape_html(text):
        if text is None:
            return ''
        s = str(text)
        return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    @classmethod
    def format_task_card(cls, task):
        tracking = cls.escape_html(task.get('trackingNumber', 'N/A'))
        l5 = cls.escape_html(task.get('l5Name', 'N/A'))
        leg = cls.escape_html(task.get('leg', 'N/A'))
        promise = cls.escape_html(task.get('promiseDate', 'N/A'))
        created = cls.escape_html(task.get('createdDate', 'N/A'))

        return "\n".join([
            f"📦 <b>Tracking Ids:</b> <code>{tracking}</code>",
            f"👤 <b>L5_Name:</b> ${l5}" if False else f"👤 <b>L5_Name:</b> {l5}",
            f"🏷 <b>Leg:</b> {leg}",
            f"📅 <b>Promise date:</b> {promise}",
            f"⏰ <b>Created date:</b> {created}"
        ])

    @classmethod
    def build_messages(cls, tasks, target_hub='MRZ'):
        if not tasks:
            return []

        total_count = len(tasks)
        cards = [cls.format_task_card(t) for t in tasks]
        card_divider = '\n────────────────────\n'
        max_length = cls.MAX_MESSAGE_LENGTH

        raw_chunks = []
        current_cards = []
        current_length = 0

        for card in cards:
            card_len_with_divider = len(card) + len(card_divider)
            if current_cards and (current_length + card_len_with_divider + 250) > max_length:
                raw_chunks.append(current_cards)
                current_cards = [card]
                current_length = len(card)
            else:
                current_cards.append(card)
                current_length += card_len_with_divider

        if current_cards:
            raw_chunks.append(current_cards)

        total_batches = len(raw_chunks)
        messages = []

        for idx, chunk_cards in enumerate(raw_chunks):
            batch_num = idx + 1
            header = f"🔔 <b>Today's Task Count -> {total_count}</b>\n"
            header += f"🏢 <b>DC:</b> {target_hub}"
            if total_batches > 1:
                header += f" | <b>Part {batch_num} of {total_batches}</b>"
            header += "\n━━━━━━━━━━━━━━━━━━━━\n"

            body = card_divider.join(chunk_cards)
            messages.append(header + body)

        return messages


# ==============================================================================
# Test Suite
# ==============================================================================

WORKBOOK_PATH = r"C:\Users\User\Desktop\res\DAILY TASK VIEW.xlsx"


class TestDailyTaskAlertPipeline(unittest.TestCase):
    """
    Rigorously tests the end-to-end task alert pipeline against
    C:\\Users\\User\\Desktop\\res\\DAILY TASK VIEW.xlsx.
    """

    @classmethod
    def setUpClass(cls):
        print(f"\n========================================================")
        print(f"LOADING WORKBOOK: {WORKBOOK_PATH}")
        print(f"========================================================")
        cls.assertTrue(cls, os.path.exists(WORKBOOK_PATH), f"Workbook file not found at: {WORKBOOK_PATH}")
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
        print(f"Workbook successfully opened. Available sheets: {cls.wb.sheetnames}")

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()
        print(f"Workbook closed successfully.\n")

    def test_01_dynamic_task_tab_resolution(self):
        """
        Step 2: Dynamically finds the task tab using regex /^tasks?\\b/i.
        Replicates TaskFilterService.resolveTaskSheet.
        """
        sheet_names = self.wb.sheetnames
        print(f"\n[Test 1] Testing sheet resolution across sheets: {sheet_names}")
        
        resolved_tab = TaskFilterService.resolve_task_sheet_name(sheet_names)
        print(f"[Test 1] Resolved sheet name: {repr(resolved_tab)}")

        self.assertEqual(resolved_tab, 'Tasks ')
        self.assertTrue(TaskFilterService.TAB_REGEX.search(resolved_tab.strip()))
        
        # Verify non-matching sheet list throws ValueError
        with self.assertRaises(ValueError):
            TaskFilterService.resolve_task_sheet_name(['Sheet84', 'Sheet83', 'R', 'Cx buddy availibilty'])

    def test_02_dynamic_header_mapping_and_aliases(self):
        """
        Step 3: Extracts Row 1 headers and verifies dynamic alias resolution
        for DC, Final_Tracking_Number, l5_name, Attribute, promise_date, created_date.
        Asserts ZERO hardcoded column indices.
        """
        ws = self.wb['Tasks ']
        header_row = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        print(f"\n[Test 2] Extracted {len(header_row)} headers from Row 1:")
        print(f"         {header_row[:15]} ...")

        resolver = HeaderResolver.create(header_row)

        resolved_indices = resolver.resolve_indices(TaskFilterService.REQUIRED_FIELDS)
        print(f"[Test 2] Dynamically resolved indices:")
        for k, v in resolved_indices.items():
            print(f"         - {k} -> Index {v} (Header: '{header_row[v]}')")

        # Verify all required keys are resolved to valid integer indices
        self.assertIn('dc', resolved_indices)
        self.assertIn('tracking_number', resolved_indices)
        self.assertIn('l5_name', resolved_indices)
        self.assertIn('leg', resolved_indices)
        self.assertIn('promise_date', resolved_indices)
        self.assertIn('created_date', resolved_indices)

        # Verify actual headers match expectations
        self.assertEqual(header_row[resolved_indices['dc']], 'DC')
        self.assertEqual(header_row[resolved_indices['tracking_number']], 'Final_Tracking_Number')
        self.assertEqual(header_row[resolved_indices['l5_name']], 'l5_name')
        self.assertEqual(header_row[resolved_indices['leg']], 'Attribute')
        self.assertEqual(header_row[resolved_indices['promise_date']], 'promise_date')
        self.assertEqual(header_row[resolved_indices['created_date']], 'created_date')

        # Test alias fallback flexibility across all required aliases
        # DC aliases: ['dc', 'hub', 'delivery_center', 'dc_name']
        self.assertEqual(resolver.get_index(['dc', 'hub', 'delivery_center']), resolved_indices['dc'])
        self.assertEqual(resolver.get_index(['hub', 'dc']), resolved_indices['dc'])
        self.assertEqual(resolver.get_index(['delivery_center', 'dc_name', 'dc']), resolved_indices['dc'])

        # Final_Tracking_Number aliases: ['final_tracking_number', 'tracking_number', 'tracking id', 'tracking']
        self.assertEqual(resolver.get_index(['tracking', 'final_tracking_number']), resolved_indices['tracking_number'])
        self.assertEqual(resolver.get_index(['tracking_number', 'final_tracking_number']), resolved_indices['tracking_number'])
        self.assertEqual(resolver.get_index('final_tracking_number'), resolved_indices['tracking_number'])

        # l5_name aliases: ['l5_name', 'l5', 'l5 name']
        self.assertEqual(resolver.get_index(['l5', 'l5_name']), resolved_indices['l5_name'])
        self.assertEqual(resolver.get_index(['l5 name', 'l5_name']), resolved_indices['l5_name'])

        # Attribute aliases: ['attribute', 'leg', 'task_leg']
        self.assertEqual(resolver.get_index(['leg', 'attribute']), resolved_indices['leg'])
        self.assertEqual(resolver.get_index(['task_leg', 'attribute']), resolved_indices['leg'])

        # promise_date aliases: ['promise_date', 'promise date']
        self.assertEqual(resolver.get_index(['promise date', 'promise_date']), resolved_indices['promise_date'])

        # created_date aliases: ['created_date', 'created date']
        self.assertEqual(resolver.get_index(['created date', 'created_date']), resolved_indices['created_date'])

        # Test missing column exception
        with self.assertRaises(KeyError):
            resolver.get_index(['non_existent_column_xyz'])

    def test_03_mrz_row_filtering(self):
        """
        Step 4: Filters rows for DC == 'MRZ' (case-insensitive, trimmed).
        Asserts that MYSP1467166924 is matched and extracted.
        """
        ws = self.wb['Tasks ']
        header_row = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        resolver = HeaderResolver.create(header_row)
        indices = resolver.resolve_indices(TaskFilterService.REQUIRED_FIELDS)

        target_hub = TaskFilterService.TARGET_HUB.strip().upper()
        filtered_tasks = []

        total_scanned = 0
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            total_scanned += 1
            raw_dc = row[indices['dc']]
            if raw_dc is None:
                continue

            dc_val = str(raw_dc).strip().upper()
            if dc_val == target_hub:
                tracking = str(row[indices['tracking_number']] or '').strip()
                l5 = str(row[indices['l5_name']] or '').strip()
                leg = str(row[indices['leg']] or '').strip()
                promise = DateUtils.format_fast(row[indices['promise_date']])
                created = DateUtils.format_fast(row[indices['created_date']])

                filtered_tasks.append({
                    'trackingNumber': tracking or 'N/A',
                    'l5Name': l5 or 'N/A',
                    'leg': leg or 'N/A',
                    'promiseDate': promise,
                    'createdDate': created,
                    'rawRowNumber': row_num
                })

        print(f"\n[Test 3] Scanned {total_scanned} data rows.")
        print(f"[Test 3] Filtered {len(filtered_tasks)} task(s) for DC '{target_hub}':")
        for t in filtered_tasks:
            print(f"         Row {t['rawRowNumber']}: Tracking={t['trackingNumber']}, L5={t['l5Name']}, Leg={t['leg']}, Promise={t['promiseDate']}, Created={t['createdDate']}")

        self.assertGreaterEqual(len(filtered_tasks), 1, "Expected at least 1 task for DC MRZ")
        mrz_task = filtered_tasks[0]
        self.assertEqual(mrz_task['trackingNumber'], 'MYSP1467166924')
        self.assertEqual(mrz_task['l5Name'], 'IMD')
        self.assertEqual(mrz_task['leg'], 'Forward')
        self.assertEqual(mrz_task['rawRowNumber'], 973)

    def test_04_date_formatting(self):
        """
        Step 5: Formats dates using DateUtils.formatFast logic.
        Asserts no 'Invalid Date' and correct format (YYYY-MM-DD HH:mm:ss or YYYY-MM-DD).
        """
        ws = self.wb['Tasks ']
        header_row = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        resolver = HeaderResolver.create(header_row)
        indices = resolver.resolve_indices(TaskFilterService.REQUIRED_FIELDS)

        # Retrieve row 973 (MRZ row)
        row_973 = None
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row_num == 973:
                row_973 = row
                break

        self.assertIsNotNone(row_973, "Row 973 could not be read.")
        raw_promise = row_973[indices['promise_date']]
        raw_created = row_973[indices['created_date']]

        formatted_promise = DateUtils.format_fast(raw_promise)
        formatted_created = DateUtils.format_fast(raw_created)

        print(f"\n[Test 4] Date formatting verification:")
        print(f"         Raw Promise: {raw_promise} -> Formatted: '{formatted_promise}'")
        print(f"         Raw Created: {raw_created} -> Formatted: '{formatted_created}'")

        self.assertNotEqual(formatted_promise, 'Invalid Date')
        self.assertNotEqual(formatted_created, 'Invalid Date')
        self.assertNotEqual(formatted_promise, 'N/A')
        self.assertNotEqual(formatted_created, 'N/A')

        date_regex = re.compile(r'^\d{4}-\d{2}-\d{2}( \d{2}:\d{2}:\d{2})?$')
        self.assertTrue(date_regex.match(formatted_promise), f"Promise date '{formatted_promise}' failed regex")
        self.assertTrue(date_regex.match(formatted_created), f"Created date '{formatted_created}' failed regex")

        self.assertEqual(formatted_promise, '2026-09-21 00:00:00')
        self.assertEqual(formatted_created, '2026-09-17 00:00:00')

        # Additional unit tests for DateUtils
        self.assertEqual(DateUtils.format_fast(None), 'N/A')
        self.assertEqual(DateUtils.format_fast(''), 'N/A')
        self.assertEqual(DateUtils.format_fast(datetime.date(2026, 9, 21)), '2026-09-21 00:00:00')

    def test_05_sha256_fingerprint_computation(self):
        """
        Step 6: Computes SHA-256 fingerprint matching GAS logic:
        tasks.map(t => f"{t.trackingNumber}|{t.l5Name}|{t.leg}|{t.promiseDate}|{t.createdDate}").join("\\n")
        Asserts 64-char hex string, determinism, and idempotency.
        """
        task = {
            'trackingNumber': 'MYSP1467166924',
            'l5Name': 'IMD',
            'leg': 'Forward',
            'promiseDate': '2026-09-21 00:00:00',
            'createdDate': '2026-09-17 00:00:00',
            'rawRowNumber': 973
        }
        tasks = [task]

        fingerprint1 = TaskFilterService.compute_fingerprint(tasks)
        fingerprint2 = TaskFilterService.compute_fingerprint(tasks)

        print(f"\n[Test 5] SHA-256 Fingerprint: {fingerprint1}")

        # Asserts valid 64-char hex string
        self.assertEqual(len(fingerprint1), 64)
        self.assertTrue(re.match(r'^[0-9a-f]{64}$', fingerprint1), "Fingerprint must be 64 lowercase hex characters")

        # Determinism
        self.assertEqual(fingerprint1, fingerprint2, "Fingerprint computation must be deterministic")

        # Exact expected digest verification
        expected_serialized = "MYSP1467166924|IMD|Forward|2026-09-21 00:00:00|2026-09-17 00:00:00"
        expected_digest = hashlib.sha256(expected_serialized.encode('utf-8')).hexdigest()
        self.assertEqual(fingerprint1, expected_digest)
        self.assertEqual(fingerprint1, "4104ea0a43d93e95c5901f2523812dcc41933fe12be4049cf15724005291b762")

        # Idempotency & mutation sensitivity
        self.assertEqual(TaskFilterService.compute_fingerprint([]), 'EMPTY_SET')
        mutated_tasks = [{**task, 'leg': 'Reverse'}]
        mutated_fp = TaskFilterService.compute_fingerprint(mutated_tasks)
        self.assertNotEqual(fingerprint1, mutated_fp)

    def test_06_telegram_message_card_rendering(self):
        """
        Step 7: Generates Telegram message cards replicating TelegramService.buildMessages:
        - Verifies HTML escaping.
        - Verifies header format:
            🔔 Today's Task Count -> {count}
            🏢 DC: MRZ
            📦 Tracking Ids: <code>MYSP1467166924</code>
            👤 L5_Name: IMD
            🏷 Leg: Forward
            📅 Promise date: 2026-09-21 00:00:00
            ⏰ Created date: 2026-09-17 00:00:00
        - Checks that total message length is strictly < 3800 characters.
        """
        # Test HTML escaping
        raw_xss = "<script>alert('XSS & Hack')</script>"
        escaped = TelegramService.escape_html(raw_xss)
        self.assertEqual(escaped, "&lt;script&gt;alert('XSS &amp; Hack')&lt;/script&gt;")

        task = {
            'trackingNumber': 'MYSP1467166924',
            'l5Name': 'IMD',
            'leg': 'Forward',
            'promiseDate': '2026-09-21 00:00:00',
            'createdDate': '2026-09-17 00:00:00'
        }
        tasks = [task]

        messages = TelegramService.build_messages(tasks, target_hub='MRZ')
        self.assertEqual(len(messages), 1)
        msg = messages[0]

        # Safe print for any terminal environment
        try:
            print(f"\n[Test 6] Rendered Telegram Message Card:\n{'-'*40}\n{msg}\n{'-'*40}")
        except UnicodeEncodeError:
            print(f"\n[Test 6] Rendered Telegram Message Card:\n{'-'*40}\n{msg.encode('ascii', 'backslashreplace').decode()}\n{'-'*40}")
        print(f"[Test 6] Message character count: {len(msg)} (limit: < 3800)")

        # Verify strict length requirement
        self.assertLess(len(msg), 3800, "Message length must be strictly < 3800 characters")

        # Verify required header lines and content
        self.assertIn("🔔 <b>Today's Task Count -> 1</b>", msg)
        self.assertIn("🏢 <b>DC:</b> MRZ", msg)
        self.assertIn("━━━━━━━━━━━━━━━━━━━━", msg)
        self.assertIn("📦 <b>Tracking Ids:</b> <code>MYSP1467166924</code>", msg)
        self.assertIn("👤 <b>L5_Name:</b> IMD", msg)
        self.assertIn("🏷 <b>Leg:</b> Forward", msg)
        self.assertIn("📅 <b>Promise date:</b> 2026-09-21 00:00:00", msg)
        self.assertIn("⏰ <b>Created date:</b> 2026-09-17 00:00:00", msg)

        # Verify batching logic under high volume (100 mock tasks)
        mock_tasks = []
        for i in range(100):
            mock_tasks.append({
                'trackingNumber': f'MYSP{1000000000 + i}',
                'l5Name': f'IMD_{i}',
                'leg': 'Forward',
                'promiseDate': '2026-09-21 00:00:00',
                'createdDate': '2026-09-17 00:00:00'
            })
        batch_messages = TelegramService.build_messages(mock_tasks, target_hub='MRZ')
        print(f"[Test 6] Batching test: 100 tasks split into {len(batch_messages)} batch(es).")
        for idx, b_msg in enumerate(batch_messages, start=1):
            self.assertLess(len(b_msg), 3800, f"Batch {idx} length {len(b_msg)} exceeds 3800 limit")
            self.assertIn(f"Part {idx} of {len(batch_messages)}", b_msg)

    def test_07_end_to_end_pipeline_simulation(self):
        """
        Full End-to-End Simulation:
        Resolves tab -> Extracts headers -> Resolves columns -> Filters MRZ -> Formats dates ->
        Computes SHA-256 fingerprint -> Renders Telegram message.
        """
        print(f"\n[Test 7] Executing Full End-to-End Pipeline Simulation...")
        # 1. Resolve Tab
        tab_name = TaskFilterService.resolve_task_sheet_name(self.wb.sheetnames)
        ws = self.wb[tab_name]

        # 2. Extract Headers & Resolve Indices
        header_row = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        resolver = HeaderResolver.create(header_row)
        indices = resolver.resolve_indices(TaskFilterService.REQUIRED_FIELDS)

        # 3. Filter Tasks
        tasks = []
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            raw_dc = row[indices['dc']]
            if raw_dc and str(raw_dc).strip().upper() == TaskFilterService.TARGET_HUB:
                tasks.append({
                    'trackingNumber': str(row[indices['tracking_number']] or '').strip(),
                    'l5Name': str(row[indices['l5_name']] or '').strip(),
                    'leg': str(row[indices['leg']] or '').strip(),
                    'promiseDate': DateUtils.format_fast(row[indices['promise_date']]),
                    'createdDate': DateUtils.format_fast(row[indices['created_date']]),
                    'rawRowNumber': row_num
                })

        # 4. Assert Filtered Task
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]['trackingNumber'], 'MYSP1467166924')

        # 5. Compute Fingerprint
        fingerprint = TaskFilterService.compute_fingerprint(tasks)
        self.assertEqual(fingerprint, '4104ea0a43d93e95c5901f2523812dcc41933fe12be4049cf15724005291b762')

        # 6. Build Messages
        messages = TelegramService.build_messages(tasks)
        self.assertEqual(len(messages), 1)
        self.assertLess(len(messages[0]), 3800)
        self.assertIn("MYSP1467166924", messages[0])
        print(f"[Test 7] Full Pipeline Simulation Passed Successfully! All assertions verified.")


if __name__ == '__main__':
    unittest.main(verbosity=2)
